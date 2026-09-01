from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from uuid import UUID

from docx import Document
from openpyxl import load_workbook

from app.procurement.strategy_exports import (
    procurement_strategy_export_csv,
    procurement_strategy_export_markdown,
    procurement_strategy_export_rows,
)
from app.schemas.projects import ProcurementStrategyView
from app.sitewise.artifact_exports import (
    render_artifact_export,
    render_table_workbook,
)


def _strategy() -> ProcurementStrategyView:
    now = datetime(2026, 8, 29, tzinfo=UTC)
    return ProcurementStrategyView.model_validate(
        {
            "id": UUID("33333333-3333-3333-3333-333333333333"),
            "project_id": UUID("22222222-2222-2222-2222-222222222222"),
            "revision": 2,
            "tenderer_column_count": 3,
            "source_fingerprint": "fingerprint",
            "created_at": now,
            "updated_at": now,
            "rows": [
                {
                    "id": UUID("44444444-4444-4444-4444-444444444444"),
                    "discipline_label": "Structural",
                    "participant_type": "consultant",
                    "request_kind": "consultant_rfp",
                    "status": "issued",
                    "notes": "",
                    "display_order": 100,
                    "origin": "derived",
                    "locked": False,
                    "candidates": [
                        {
                            "id": UUID("55555555-5555-5555-5555-555555555555"),
                            "slot": 1,
                            "company_name": "North & Co",
                        }
                    ],
                }
            ],
        }
    )


def test_procurement_strategy_uses_standard_word_and_workbook_renderers() -> None:
    strategy = _strategy()
    rows = procurement_strategy_export_rows(strategy)

    docx = render_artifact_export(
        procurement_strategy_export_markdown(strategy),
        export_format="docx",
        project_title="Demo Project",
        artifact_title="Procurement Strategy",
        version=strategy.revision,
    )
    document = Document(BytesIO(docx))
    assert document.core_properties.title == "Procurement Strategy"
    assert document.tables[0].cell(0, 0).text == "Discipline"
    assert document.tables[0].cell(1, 1).text == "North & Co"

    workbook = load_workbook(
        BytesIO(
            render_table_workbook(
                project_title="Demo Project",
                artifact_title="Procurement Strategy",
                version=strategy.revision,
                sheet_title="Procurement",
                headers=rows[0],
                rows=rows[1:],
            )
        )
    )
    worksheet = workbook["Procurement"]
    assert worksheet["A1"].value == "Procurement Strategy"
    assert worksheet["A4"].value == "Discipline"
    assert worksheet["B5"].value == "North & Co"
    assert worksheet["A4"].fill.fgColor.rgb == "002C3037"
    assert procurement_strategy_export_csv(strategy).splitlines() == [
        "Discipline,Firm 1,Firm 2,Firm 3,Status",
        "Structural,North & Co,,,Issued",
    ]
