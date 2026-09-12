from __future__ import annotations

import uuid
from datetime import UTC, datetime
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

from docx import Document
from fastapi.testclient import TestClient

from app.auth.dependencies import CurrentUser, get_current_user
from app.database.project import Project
from app.database.session import get_db
from app.main import fastapi_app as app
from app.projects.profile_exports import profile_export_markdown
from app.schemas.projects import ProjectAsset, ProjectProfileView, ProjectSubclassSelection
from app.sitewise.artifact_exports import render_artifact_export

USER_ID = uuid.UUID("11111111-1111-1111-1111-111111111111")
PROJECT_ID = uuid.UUID("22222222-2222-2222-2222-222222222222")
NOW = datetime(2026, 9, 7, 12, 0, tzinfo=UTC)


def _profile() -> ProjectProfileView:
    return ProjectProfileView(
        project_id=PROJECT_ID,
        profile_revision=4,
        title="Demo Project",
        building_class="commercial",
        work_type="refurb",
        subclasses=["office"],
        scale={"nla_sqm": 1200},
        complexity={"operational_constraints": "live_environment"},
        work_scope=["fire_services"],
        assets=[
            ProjectAsset(type="Chiller", location="Roof plant", action="Replace"),
        ],
        scope_narrative=["Replace the existing chiller and associated pipework."],
        budget="$1.2m",
        state="NSW",
        site_address="82 Queen Street, Petersham NSW 2049",
        client="Acme Developments",
    )


def test_profile_export_markdown_uses_taxonomy_labels_for_every_field() -> None:
    markdown = profile_export_markdown(_profile())

    assert markdown.startswith("# Project Profile")
    assert "Demo Project" in markdown
    assert "82 Queen Street, Petersham NSW 2049" in markdown
    assert "Acme Developments" in markdown
    assert "NSW" in markdown
    assert "Commercial" in markdown
    assert "Refurbishment" in markdown
    assert "Office (Class 5)" in markdown
    assert "$1.2m" in markdown
    assert "NLA sqm" in markdown
    assert "1200" in markdown
    assert "Operational constraints" in markdown
    assert "Live Environment" in markdown
    assert "Fire Services Upgrade" in markdown
    assert "Replace the existing chiller and associated pipework." in markdown
    assert "## Scope notes" in markdown
    assert "Chiller" in markdown
    assert "Roof plant" in markdown
    assert "Contamination" in markdown
    assert "Not stated" in markdown


def test_profile_export_groups_work_scope_by_profile_categories() -> None:
    profile = _profile()
    profile.work_scope = [
        "fire_services",
        "stripout",
        "partitions_walls",
        "ceilings",
    ]
    markdown = profile_export_markdown(profile)

    assert "### Investigation and Strip-Out" in markdown
    assert "### Internal Fit-Out" in markdown
    assert "### Services and Compliance Upgrade" in markdown
    assert "- Strip-Out Works" in markdown
    assert "- Partitions/Internal Walls" in markdown
    assert "- Ceilings and Acoustic Treatments" in markdown
    assert "- Fire Services Upgrade" in markdown
    investigation = markdown.index("Investigation and Strip-Out")
    fitout = markdown.index("Internal Fit-Out")
    services = markdown.index("Services and Compliance Upgrade")
    assert investigation < fitout < services


def test_profile_export_includes_every_form_field_and_selected_scope_item() -> None:
    profile = ProjectProfileView(
        project_id=PROJECT_ID,
        profile_revision=2,
        title="Seven Hills Townhouse",
        building_class="residential",
        work_type="new",
        subclasses=["townhouses"],
        scale={"site_sqm": 3240, "dwellings": 12},
        complexity={"planning": "da", "procurement_route": "traditional"},
        work_scope=[
            "demolition",
            "site_clearance",
            "bulk_earthworks",
            "utility_diversions",
            "detailed_earthworks",
            "site_drainage",
            "stormwater_management",
            "internal_roads_pavements",
            "retaining_walls",
            "substructure_foundations",
            "superstructure",
            "facade_system",
            "roofing_system",
            "glazing_windows",
            "waterproofing",
            "mechanical_hvac",
            "electrical_power",
            "lighting",
            "hydraulic_plumbing",
            "fire_services",
            "vertical_transport",
            "security_systems",
            "partitions_internal_walls",
            "ceilings",
            "flooring",
            "joinery_cabinetry",
            "specialist_fitout",
            "landscaping",
            "car_parking",
            "signage_wayfinding",
            "external_lighting",
            "fencing_gates",
        ],
        scope_narrative=[
            "Twelve attached two-storey dwellings at feasibility: eight 3-bedroom and four 4-bedroom",
        ],
        budget="$9,800,000 excl GST",
        state="NSW",
        site_address="14–18 Wianamatta Avenue, Seven Hills NSW 2147",
        client="Wianamatta Developments Pty Ltd",
    )
    markdown = profile_export_markdown(profile)
    document = Document(
        BytesIO(
            render_artifact_export(
                markdown,
                export_format="docx",
                project_title="Seven Hills Townhouse",
                artifact_title="Project Profile",
                version=2,
            )
        )
    )
    text = "\n".join(
        [
            *(paragraph.text for paragraph in document.paragraphs),
            *(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ),
        ]
    )

    assert "Site sqm" in markdown
    assert "3240" in markdown
    assert "Dwellings" in markdown
    assert "GFA sqm" in markdown
    assert "Average dwelling size sqm" in markdown
    assert "Planning" in markdown
    assert "DA" in markdown
    assert "Procurement route" in markdown
    assert "Traditional (Lump Sum)" in markdown
    assert "Contamination level" in markdown
    assert "Access constraints" in markdown
    assert "Heritage status" in markdown
    assert "Twelve attached two-storey dwellings" in markdown
    for item in (
        "Demolition",
        "Site Clearance",
        "Stormwater Management",
        "Facade System",
        "Specialist Fitout (Lab/Kitchen)",
        "Fencing/Gates",
    ):
        assert item in markdown
        assert item in text
    assert "Wianamatta Developments Pty Ltd" in text
    assert "GFA sqm" in text
    assert "Heritage status" in text
    assert "Not stated" in text


def test_profile_export_markdown_uses_custom_other_subclass_label() -> None:
    profile = _profile()
    profile.subclasses = [
        ProjectSubclassSelection(value="other", label="Laboratory office"),
    ]

    markdown = profile_export_markdown(profile)

    assert "Laboratory office" in markdown


def test_profile_export_word_contains_the_labelled_profile() -> None:
    document = Document(
        BytesIO(
            render_artifact_export(
                profile_export_markdown(_profile()),
                export_format="docx",
                project_title="Demo Project",
                artifact_title="Project Profile",
                version=4,
            )
        )
    )
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )

    assert "Project Profile" in text or "Project Profile" in table_text
    assert "Acme Developments" in text or "Acme Developments" in table_text
    assert "Commercial" in text or "Commercial" in table_text
    assert "Office (Class 5)" in text or "Office (Class 5)" in table_text


def test_profile_export_word_keeps_every_work_scope_item() -> None:
    profile = _profile()
    profile.work_scope = ["fire_services", "stripout", "partitions_walls"]
    document = Document(
        BytesIO(
            render_artifact_export(
                profile_export_markdown(profile),
                export_format="docx",
                project_title="Demo Project",
                artifact_title="Project Profile",
                version=4,
            )
        )
    )
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "Investigation and Strip-Out" in text
    assert "Internal Fit-Out" in text
    assert "Services and Compliance Upgrade" in text
    assert "Strip-Out Works" in text
    assert "Fire Services Upgrade" in text
    assert "Partitions/Internal Walls" in text


def _project() -> Project:
    return Project(
        id=PROJECT_ID,
        owner_user_id=USER_ID,
        slug="demo",
        title="Demo Project",
        workspace_path="04-projects/demo",
        phase="brief-planning",
        archetype="small-commercial",
        building_class="commercial",
        work_type="refurb",
        user_role="architect-pm",
        state="NSW",
        status="active",
        profile_revision=4,
        project_metadata={
            "taxonomy": {
                "subclasses": ["office"],
                "scale": {"nla_sqm": 1200},
                "complexity": {"operational_constraints": "live_environment"},
                "work_scope": ["fire_services"],
                "budget": "$1.2m",
                "site_address": "82 Queen Street, Petersham NSW 2049",
                "client": "Acme Developments",
                "scope_narrative": ["Replace the existing chiller and associated pipework."],
            }
        },
        created_at=NOW,
        updated_at=NOW,
    )


@pytest.fixture
def client() -> TestClient:
    current_user = CurrentUser(id=USER_ID, email="user@example.com")

    async def override_get_db():
        yield AsyncMock()

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = lambda: current_user
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def test_export_project_profile_renders_word(client: TestClient) -> None:
    render_export = MagicMock(return_value=b"PK word-bytes")

    with (
        patch("app.api.projects.get_project", new=AsyncMock(return_value=_project())),
        patch("app.api.projects.require_active_entitlement", new=AsyncMock()),
        patch("app.api.projects.render_artifact_export", new=render_export),
    ):
        response = client.get(
            f"/projects/{PROJECT_ID}/profile/export",
            params={"format": "docx"},
        )

    assert response.status_code == 200
    assert response.content == b"PK word-bytes"
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert 'filename="Project_Profile_v04.docx"' in response.headers["content-disposition"]
    render_export.assert_called_once()
    kwargs = render_export.call_args.kwargs
    assert kwargs["export_format"] == "docx"
    assert kwargs["project_title"] == "Demo Project"
    assert kwargs["artifact_title"] == "Project Profile"
    assert kwargs["version"] == 4
    assert "Commercial" in render_export.call_args.args[0]
    assert "Acme Developments" in render_export.call_args.args[0]


def test_export_project_profile_returns_markdown(client: TestClient) -> None:
    with (
        patch("app.api.projects.get_project", new=AsyncMock(return_value=_project())),
        patch("app.api.projects.require_active_entitlement", new=AsyncMock()),
    ):
        response = client.get(
            f"/projects/{PROJECT_ID}/profile/export",
            params={"format": "md"},
        )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/markdown")
    assert 'filename="Project_Profile_v04.md"' in response.headers["content-disposition"]
    assert "# Project Profile" in response.text
    assert "Acme Developments" in response.text
    assert "Commercial" in response.text


def test_export_project_profile_renders_pdf(client: TestClient) -> None:
    with (
        patch("app.api.projects.get_project", new=AsyncMock(return_value=_project())),
        patch("app.api.projects.require_active_entitlement", new=AsyncMock()),
        patch(
            "app.api.projects.render_artifact_export",
            new=MagicMock(return_value=b"%PDF-1.7 profile"),
        ),
    ):
        response = client.get(
            f"/projects/{PROJECT_ID}/profile/export",
            params={"format": "pdf"},
        )

    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    assert response.headers["content-type"] == "application/pdf"
    assert 'filename="Project_Profile_v04.pdf"' in response.headers["content-disposition"]
